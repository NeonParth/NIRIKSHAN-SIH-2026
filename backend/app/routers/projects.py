from __future__ import annotations

import io
from uuid import UUID

from docx import Document
from fastapi import APIRouter, Depends, Query, Response, status
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.database.connection import get_db
from app.database.models import Project, RiskAssessment
from app.database.models.enums import CompletionStatus, DataSource, RiskLevel
from app.database.schemas.api import Paginated, ProjectCreate, ProjectDetail, ProjectSummary
from app.services.assessment_service import to_project_detail, to_project_summary
from app.services.gis_service import point_wkt
from app.utils.errors import http_error
from app.utils.pagination import page_offset, pagination_params
from geoalchemy2.elements import WKTElement
from geoalchemy2.types import Geography
from sqlalchemy import cast

router = APIRouter(prefix="/projects", tags=["projects"])


def _base_query():
    return select(Project).options(selectinload(Project.risk_assessments), selectinload(Project.evidence))


@router.get("", response_model=Paginated[ProjectSummary])
def list_projects(
    q: str | None = Query(default=None, description="Search name or code"),
    data_source: DataSource | None = None,
    completion_status: CompletionStatus | None = None,
    risk_level: RiskLevel | None = None,
    longitude: float | None = Query(default=None, ge=-180, le=180),
    latitude: float | None = Query(default=None, ge=-90, le=90),
    radius_m: float | None = Query(default=None, gt=0),
    paging: tuple[int, int] = Depends(pagination_params),
    db: Session = Depends(get_db),
):
    page, page_size = paging
    stmt = _base_query()
    if q:
        pattern = f"%{q}%"
        stmt = stmt.where(or_(Project.name.ilike(pattern), Project.code.ilike(pattern)))
    if data_source:
        stmt = stmt.where(Project.data_source == data_source)
    if completion_status:
        stmt = stmt.where(Project.completion_status == completion_status)
    if risk_level:
        latest = (
            select(RiskAssessment.project_id, func.max(RiskAssessment.created_at).label("latest_at"))
            .group_by(RiskAssessment.project_id)
            .subquery()
        )
        stmt = (
            stmt.join(latest, latest.c.project_id == Project.id)
            .join(
                RiskAssessment,
                (RiskAssessment.project_id == Project.id) & (RiskAssessment.created_at == latest.c.latest_at),
            )
            .where(RiskAssessment.risk_level == risk_level)
        )
    if longitude is not None and latitude is not None and radius_m is not None:
        origin = WKTElement(point_wkt(longitude, latitude), srid=4326)
        stmt = stmt.where(Project.location.is_not(None)).where(
            func.ST_DWithin(cast(Project.location, Geography()), cast(origin, Geography()), radius_m)
        )
    elif any(value is not None for value in (longitude, latitude, radius_m)):
        raise http_error(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "INVALID_GEO_FILTER",
            "longitude, latitude, and radius_m must be provided together",
        )

    count_stmt = select(func.count()).select_from(stmt.order_by(None).subquery())
    total = db.scalar(count_stmt) or 0
    rows = db.scalars(stmt.order_by(Project.created_at.desc()).offset(page_offset(page, page_size)).limit(page_size)).all()
    return Paginated(
        items=[to_project_summary(item) for item in rows],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/{project_id}", response_model=ProjectDetail)
def get_project(project_id: UUID, db: Session = Depends(get_db)):
    project = db.scalars(_base_query().where(Project.id == project_id)).first()
    if project is None:
        raise http_error(status.HTTP_404_NOT_FOUND, "NOT_FOUND", "Project not found")
    return to_project_detail(project)


@router.get("/{project_id}/investigation-report.docx")
def download_investigation_report(project_id: UUID, db: Session = Depends(get_db)):
    project = db.scalars(_base_query().where(Project.id == project_id)).first()
    if project is None:
        raise http_error(status.HTTP_404_NOT_FOUND, "NOT_FOUND", "Project not found")

    detail = to_project_detail(project)

    doc = Document()
    doc.add_heading(f"NIRIKSHAN Project Investigation Report - {detail.code}", level=0)

    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph(f"Project Name: {detail.name}")
    doc.add_paragraph(f"Project Code: {detail.code}")
    doc.add_paragraph(f"Completion Status: {detail.completion_status}")
    doc.add_paragraph(f"Data Source: {detail.data_source}")
    if detail.allocated_amount is not None:
        doc.add_paragraph(f"Allocated Amount: INR {detail.allocated_amount:,.2f}")
    if detail.location:
        doc.add_paragraph(f"Location: Latitude {detail.location.latitude}, Longitude {detail.location.longitude}")

    doc.add_heading("Risk Assessment", level=1)
    if detail.latest_assessment:
        assessment = detail.latest_assessment
        doc.add_paragraph(f"Risk Score: {assessment.risk_score}/100")
        doc.add_paragraph(f"Risk Level: {assessment.risk_level}")
        doc.add_paragraph(f"Verification Status: {assessment.verification_status}")
        doc.add_paragraph(f"Analysis Type: {assessment.analysis_type}")
        if assessment.reviewer_remarks:
            doc.add_paragraph(f"Reviewer Remarks: {assessment.reviewer_remarks}")

        doc.add_heading("Contributing Risk Signals", level=2)
        if assessment.contributing_signals:
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Signal"
            hdr_cells[1].text = "Contribution"
            hdr_cells[2].text = "Explanation"
            for sig in assessment.contributing_signals:
                row_cells = table.add_row().cells
                row_cells[0].text = sig.signal_name
                row_cells[1].text = f"+{sig.contribution}"
                row_cells[2].text = sig.explanation
        else:
            doc.add_paragraph("No contributing signals recorded.")

        doc.add_heading("Evidence References", level=2)
        if assessment.evidence_references:
            for ref in assessment.evidence_references:
                doc.add_paragraph(f"• Evidence ID: {ref.evidence_id} - Relevance: {ref.relevance}\n  {ref.explanation}")
        else:
            doc.add_paragraph("No evidence references recorded.")
    else:
        doc.add_paragraph("No risk assessment recorded for this project.")

    doc.add_heading("Disclaimer", level=1)
    doc.add_paragraph(
        "Demo data only. SYNTHETIC records are not official MPLADS data. "
        "Risk scores indicate potential anomalies for verification—not fraud."
    )

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    filename = f"investigation-report-{detail.code}.docx"
    return Response(
        content=file_stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



@router.post("", response_model=ProjectDetail, status_code=status.HTTP_201_CREATED)
def create_project(payload: ProjectCreate, db: Session = Depends(get_db)):
    existing = db.scalar(select(Project).where(Project.code == payload.code))
    if existing:
        raise http_error(status.HTTP_409_CONFLICT, "CONFLICT", "Project code already exists")
    location = None
    if payload.longitude is not None and payload.latitude is not None:
        location = WKTElement(point_wkt(payload.longitude, payload.latitude), srid=4326)
    elif payload.longitude is not None or payload.latitude is not None:
        raise http_error(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "INVALID_LOCATION",
            "longitude and latitude must be provided together",
        )
    project = Project(
        code=payload.code,
        name=payload.name,
        description=payload.description,
        location=location,
        allocated_amount=payload.allocated_amount,
        benchmark_amount=payload.benchmark_amount,
        sanctioned_date=payload.sanctioned_date,
        completion_status=payload.completion_status,
        data_source=payload.data_source,
    )
    db.add(project)
    db.commit()
    loaded = db.scalars(_base_query().where(Project.id == project.id)).one()
    return to_project_detail(loaded)
